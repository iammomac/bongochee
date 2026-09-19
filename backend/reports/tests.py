import io
from datetime import date, timedelta

import openpyxl
from django.utils import timezone
from rest_framework import status
from rest_framework.test import APITestCase

from accounts.models import PasswordChangeRequest, User
from catalog.models import Category, PhoneModel
from rbac.models import Permission, Role
from returns_app.models import Return
from sales.models import Sale, SaleItem
from stock.models import StockIn, StockItem
from suppliers.models import Supplier


class DashboardSummaryTests(APITestCase):
    def setUp(self):
        perm = Permission.objects.create(codename="view_dashboard", label="View Dashboard", category="general")
        role = Role.objects.create(name="Viewer")
        role.permissions.add(perm)
        self.user = User.objects.create_user(
            username="viewer",
            password="Str0ngPassw0rd!",
            phone="255700000021",
            role=role,
            must_change_password=False,
        )
        self.client.force_authenticate(self.user)

        category = Category.objects.create(name="Samsung")
        model = PhoneModel.objects.create(category=category, name="Galaxy A56")
        supplier = Supplier.objects.create(name="Blue Telecom")
        stock_in = StockIn.objects.create(supplier=supplier, import_date=date.today(), created_by=self.user)
        stock_item = StockItem.objects.create(
            stock_in=stock_in,
            category=category,
            model=model,
            quantity=5,
            quantity_remaining=5,
            buying_price="500000",
            min_selling_price="600000",
            max_selling_price="700000",
        )
        sale = Sale.objects.create(
            invoice_number="INV-1", customer_name="Walk-in", payment_method="cash", sold_by=self.user
        )
        self.sale_item = SaleItem.objects.create(
            sale=sale, stock_item=stock_item, imei="123456789012345", selling_price="650000", discount="50000"
        )

    def test_revenue_trend_nets_out_discount(self):
        res = self.client.get("/api/v1/reports/dashboard-summary/")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        trend = res.json()["revenueTrend"]
        self.assertEqual(len(trend), 14)
        dates = [point["date"] for point in trend]
        self.assertEqual(dates, sorted(dates))
        today_point = trend[-1]
        self.assertEqual(today_point["date"], date.today().isoformat())
        # 650000 sold - 50000 discount = 600000 net revenue for the day.
        self.assertEqual(float(today_point["revenue"]), 600000.0)

    def test_todays_profit_nets_out_discount(self):
        res = self.client.get("/api/v1/reports/dashboard-summary/")
        # net revenue 600000 - buying_price 500000 = 100000 profit.
        self.assertEqual(float(res.json()["todaysProfit"]), 100000.0)

    def test_todays_returns_reflects_real_data(self):
        Return.objects.create(
            sale_item=self.sale_item,
            return_date=date.today(),
            return_category="battery",
            processed_by=self.user,
        )
        res = self.client.get("/api/v1/reports/dashboard-summary/")
        self.assertEqual(res.json()["todaysReturns"], 1)

    def test_pending_returns_counts_pending_and_processing_only(self):
        resolved = Return.objects.create(
            sale_item=self.sale_item, return_date=date.today(), return_category="battery",
            status="resolved", processed_by=self.user,
        )
        Return.objects.create(
            sale_item=self.sale_item, return_date=date.today(), return_category="camera",
            status="pending", processed_by=self.user,
        )
        res = self.client.get("/api/v1/reports/dashboard-summary/")
        self.assertEqual(res.json()["pendingReturns"], 1)
        resolved.delete()

    def test_pending_password_requests_counts_pending_only(self):
        PasswordChangeRequest.objects.create(user=self.user, status="pending")
        PasswordChangeRequest.objects.create(user=self.user, status="approved")
        res = self.client.get("/api/v1/reports/dashboard-summary/")
        self.assertEqual(res.json()["pendingPasswordRequests"], 1)


class ReportFixtures(APITestCase):
    """One sale of three units of one model (one profitable, one below buying price,
    one below minimum), a return on the first, and three users with different
    report permissions -- shared by the endpoint and detail-table tests."""

    def setUp(self):
        self.view_reports = Permission.objects.create(codename="view_reports", label="View Reports", category="reports")
        self.export_reports = Permission.objects.create(
            codename="export_reports", label="Export Reports", category="reports"
        )
        self.view_profit = Permission.objects.create(codename="view_profit", label="View Profit", category="reports")
        self.manage_users = Permission.objects.create(codename="manage_users", label="Manage Users", category="admin")

        viewer_role = Role.objects.create(name="ReportViewer")
        viewer_role.permissions.add(self.view_reports)
        self.viewer = User.objects.create_user(
            username="reportviewer",
            password="Str0ngPassw0rd!",
            phone="255700000070",
            role=viewer_role,
            must_change_password=False,
        )

        full_role = Role.objects.create(name="ReportFull")
        full_role.permissions.add(self.view_reports, self.export_reports, self.view_profit, self.manage_users)
        self.full_user = User.objects.create_user(
            username="reportfull",
            password="Str0ngPassw0rd!",
            phone="255700000071",
            role=full_role,
            must_change_password=False,
        )

        salesperson_role = Role.objects.create(name="Salesperson")
        self.salesperson = User.objects.create_user(
            username="salesperson1",
            password="Str0ngPassw0rd!",
            phone="255700000072",
            role=salesperson_role,
            must_change_password=False,
        )

        self.supplier = Supplier.objects.create(name="Blue Telecom")
        self.category = Category.objects.create(name="Samsung")
        self.model = PhoneModel.objects.create(category=self.category, name="Galaxy A56")
        self.stock_in = StockIn.objects.create(
            supplier=self.supplier, import_date=date.today(), created_by=self.salesperson
        )
        self.stock_item = StockItem.objects.create(
            stock_in=self.stock_in,
            category=self.category,
            model=self.model,
            quantity=10,
            quantity_remaining=7,
            buying_price="500000",
            min_selling_price="600000",
            max_selling_price="700000",
        )

        sale = Sale.objects.create(
            invoice_number="INV-REPORT-1", customer_name="Walk-in", payment_method="cash", sold_by=self.salesperson
        )
        # Profitable: net 650000, well above both buying (500000) and min (600000).
        self.profitable_item = SaleItem.objects.create(
            sale=sale, stock_item=self.stock_item, imei="100000000000001", selling_price="650000", discount="0"
        )
        # Below buying price: net 480000 < buying_price 500000 — a true loss.
        self.below_buying_item = SaleItem.objects.create(
            sale=sale, stock_item=self.stock_item, imei="100000000000002", selling_price="480000", discount="0"
        )
        # Below minimum price only: net 550000 — above buying (500000) but below min (600000).
        self.below_minimum_item = SaleItem.objects.create(
            sale=sale, stock_item=self.stock_item, imei="100000000000003", selling_price="650000", discount="100000"
        )

        Return.objects.create(
            sale_item=self.profitable_item,
            return_date=date.today(),
            return_category="battery",
            processed_by=self.salesperson,
        )

        self.client.force_authenticate(self.full_user)

    def _range(self):
        today = date.today().isoformat()
        return {"date_from": today, "date_to": today}


class ReportEndpointsTests(ReportFixtures):
    def test_sales_summary_totals(self):
        res = self.client.get("/api/v1/reports/sales-summary/", self._range())
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        totals = res.json()["totals"]
        self.assertEqual(totals["units"], 3)
        # net revenue = 650000 + 480000 + 550000 = 1680000
        self.assertEqual(float(totals["revenue"]), 1680000.0)

    def test_sales_summary_group_by_category(self):
        res = self.client.get("/api/v1/reports/sales-summary/", {**self._range(), "group_by": "category"})
        rows = res.json()["rows"]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["label"], "Samsung")
        self.assertEqual(rows[0]["units"], 3)

    def test_loss_report_classifies_both_types(self):
        res = self.client.get("/api/v1/reports/loss/", self._range())
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        rows = res.json()["rows"]
        self.assertEqual(len(rows), 2)
        types = {row["imei"]: row["lossTypeDisplay"] for row in rows}
        self.assertEqual(types["100000000000002"], "Below buying price")
        self.assertEqual(types["100000000000003"], "Below minimum price")

    def test_stock_summary_by_category(self):
        res = self.client.get("/api/v1/reports/stock-summary/", {"group_by": "category"})
        rows = res.json()["rows"]
        self.assertEqual(rows[0]["label"], "Samsung")
        self.assertEqual(rows[0]["quantity"], 7)

    def test_supplier_summary(self):
        res = self.client.get("/api/v1/reports/supplier-summary/", self._range())
        rows = res.json()["rows"]
        self.assertEqual(rows[0]["label"], "Blue Telecom")
        self.assertEqual(rows[0]["quantity"], 10)

    def test_returns_summary_by_category(self):
        res = self.client.get("/api/v1/reports/returns-summary/", self._range())
        rows = res.json()["rows"]
        self.assertEqual(rows[0]["count"], 1)

    def test_person_report_shape(self):
        res = self.client.get(f"/api/v1/reports/person/{self.salesperson.id}/", self._range())
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        body = res.json()
        self.assertEqual(body["stockAdded"]["batches"], 1)
        self.assertEqual(body["salesMade"]["units"], 3)
        self.assertEqual(body["returnsProcessed"], 1)

    def test_export_requires_export_reports_permission(self):
        self.client.force_authenticate(self.viewer)
        res = self.client.get("/api/v1/reports/sales-summary/", {**self._range(), "export": "xlsx"})
        self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)

        res = self.client.get("/api/v1/reports/sales-summary/", self._range())
        self.assertEqual(res.status_code, status.HTTP_200_OK)

    def test_export_xlsx_returns_file(self):
        res = self.client.get("/api/v1/reports/sales-summary/", {**self._range(), "export": "xlsx"})
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertIn("spreadsheetml", res["Content-Type"])

    def test_export_pdf_returns_file(self):
        res = self.client.get("/api/v1/reports/sales-summary/", {**self._range(), "export": "pdf"})
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res["Content-Type"], "application/pdf")

    def test_view_profit_gates_profit_data(self):
        self.client.force_authenticate(self.viewer)  # has view_reports but not view_profit
        res = self.client.get("/api/v1/reports/sales-summary/", self._range())
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertNotIn("profit", res.json()["totals"])

        res = self.client.get("/api/v1/reports/loss/", self._range())
        self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)


class ReportDetailTests(ReportFixtures):
    """Every report type carries the full per-transaction table, not just the
    grouped summary -- see BaseReportView.build_details."""

    def setUp(self):
        super().setUp()
        self.stock_item.notes = "Full box"
        self.stock_item.save(update_fields=["notes"])
        Sale.objects.filter(invoice_number="INV-REPORT-1").update(notes="Pays balance Friday")

    def test_sales_detail_row_carries_every_requested_field(self):
        res = self.client.get("/api/v1/reports/sales-summary/", self._range())
        details = res.json()["details"]
        # Three units of the same model in one sale collapse into a single row.
        self.assertEqual(len(details), 1)
        row = details[0]
        self.assertRegex(row["time"], r"^\d{2}:\d{2}$")
        self.assertRegex(row["date"], r"^\d{4}-\d{2}-\d{2}$")
        self.assertEqual(row["invoiceNumber"], "INV-REPORT-1")
        self.assertEqual(row["soldByName"], "salesperson1")
        self.assertEqual(row["customerName"], "Walk-in")
        self.assertEqual(row["categoryName"], "Samsung")
        self.assertEqual(row["modelName"], "Galaxy A56")
        self.assertEqual(row["supplierName"], "Blue Telecom")
        self.assertEqual(row["units"], 3)
        self.assertEqual(float(row["priceSold"]), 1780000.0)  # 650000 + 480000 + 650000
        self.assertEqual(float(row["discount"]), 100000.0)
        self.assertEqual(float(row["revenue"]), 1680000.0)
        self.assertEqual(float(row["profit"]), 180000.0)  # 1680000 - 3 x 500000
        self.assertEqual(row["condition"], "Full box")
        self.assertEqual(row["saleNotes"], "Pays balance Friday")

    def test_sales_detail_totals_row_up_the_columns(self):
        totals = self.client.get("/api/v1/reports/sales-summary/", self._range()).json()["detailTotals"]
        self.assertEqual(totals["units"], 3)
        self.assertEqual(float(totals["revenue"]), 1680000.0)
        self.assertEqual(float(totals["profit"]), 180000.0)

    def test_sales_detail_is_the_same_whatever_the_group_by(self):
        for group_by in ("day", "category", "model", "user", "payment_method", "supplier"):
            res = self.client.get("/api/v1/reports/sales-summary/", {**self._range(), "group_by": group_by})
            self.assertEqual(len(res.json()["details"]), 1, group_by)

    def test_sales_detail_respects_filters(self):
        other = Category.objects.create(name="Apple")
        res = self.client.get("/api/v1/reports/sales-summary/", {**self._range(), "category": str(other.id)})
        self.assertEqual(res.json()["details"], [])

    def test_detail_profit_is_hidden_without_view_profit(self):
        self.client.force_authenticate(self.viewer)  # view_reports only
        for path in ("sales-summary", "returns-summary", "stock-summary", "supplier-summary"):
            body = self.client.get(f"/api/v1/reports/{path}/", self._range()).json()
            self.assertTrue(body["details"], path)
            for row in body["details"]:
                self.assertNotIn("profit", row, path)
            self.assertNotIn("profit", body["detailTotals"], path)
        stock = self.client.get("/api/v1/reports/stock-summary/").json()["details"][0]
        self.assertNotIn("buyingPrice", stock)
        self.assertNotIn("stockValue", stock)

    def test_returns_detail_row(self):
        row = self.client.get("/api/v1/reports/returns-summary/", self._range()).json()["details"][0]
        self.assertRegex(row["time"], r"^\d{2}:\d{2}$")
        self.assertEqual(row["invoiceNumber"], "INV-REPORT-1")
        self.assertEqual(row["processedByName"], "salesperson1")
        self.assertEqual(row["customerName"], "Walk-in")
        self.assertEqual(row["supplierName"], "Blue Telecom")
        self.assertEqual(row["units"], 1)
        self.assertEqual(float(row["priceSold"]), 650000.0)
        self.assertEqual(float(row["revenue"]), 650000.0)
        self.assertEqual(float(row["profit"]), 150000.0)
        self.assertEqual(row["condition"], "Full box")
        self.assertEqual(row["issue"], "Battery")
        self.assertEqual(row["status"], "Pending")

    def test_stock_detail_row_shows_what_sold_from_the_line(self):
        row = self.client.get("/api/v1/reports/stock-summary/").json()["details"][0]
        self.assertEqual(row["addedByName"], "salesperson1")
        self.assertEqual(row["supplierName"], "Blue Telecom")
        self.assertEqual(row["quantity"], 10)
        self.assertEqual(row["quantityRemaining"], 7)
        self.assertEqual(row["unitsSold"], 3)
        self.assertEqual(float(row["stockValue"]), 3500000.0)  # 7 x 500000
        self.assertEqual(float(row["revenue"]), 1680000.0)
        self.assertEqual(float(row["profit"]), 180000.0)
        self.assertEqual(row["condition"], "Full box")

    def test_supplier_detail_is_limited_to_the_import_date_range(self):
        inside = self.client.get("/api/v1/reports/supplier-summary/", self._range()).json()["details"]
        self.assertEqual(len(inside), 1)
        outside = self.client.get(
            "/api/v1/reports/supplier-summary/", {"date_from": "2001-01-01", "date_to": "2001-01-02"}
        ).json()["details"]
        self.assertEqual(outside, [])

    def test_loss_rows_carry_the_same_fields(self):
        row = self.client.get("/api/v1/reports/loss/", self._range()).json()["rows"][0]
        for key in (
            "time", "soldByName", "supplierName", "units", "priceSold", "discount", "profit", "condition", "saleNotes",
        ):
            self.assertIn(key, row)
        self.assertEqual(row["soldByName"], "salesperson1")
        self.assertEqual(row["condition"], "Full box")

    def test_xlsx_export_has_a_summary_and_a_details_sheet(self):
        res = self.client.get("/api/v1/reports/sales-summary/", {**self._range(), "export": "xlsx"})
        workbook = openpyxl.load_workbook(io.BytesIO(res.content))
        self.assertEqual(workbook.sheetnames, ["Summary", "Details"])
        details = list(workbook["Details"].iter_rows(values_only=True))
        self.assertEqual(details[0][:4], ("Date", "Time", "Invoice", "Salesperson"))
        self.assertEqual(details[1][2], "INV-REPORT-1")
        self.assertEqual(details[-1][0], "Total")  # grand-total row under the detail lines
        self.assertIn("Condition", details[0])
        self.assertIn("Full box", details[1])

    def test_pdf_export_with_details_is_a_valid_pdf(self):
        for path in ("sales-summary", "returns-summary", "stock-summary", "supplier-summary", "loss"):
            res = self.client.get(f"/api/v1/reports/{path}/", {**self._range(), "export": "pdf"})
            self.assertEqual(res.status_code, status.HTTP_200_OK, path)
            self.assertTrue(res.content.startswith(b"%PDF"), path)

    def test_xlsx_export_without_view_profit_omits_profit_columns(self):
        exporter_role = Role.objects.create(name="ExportOnly")
        exporter_role.permissions.add(self.view_reports, self.export_reports)
        exporter = User.objects.create_user(
            username="exportonly", password="Str0ngPassw0rd!", phone="255700000073",
            role=exporter_role, must_change_password=False,
        )
        self.client.force_authenticate(exporter)
        res = self.client.get("/api/v1/reports/sales-summary/", {**self._range(), "export": "xlsx"})
        workbook = openpyxl.load_workbook(io.BytesIO(res.content))
        header = [cell.value for cell in workbook["Details"][1]]
        self.assertNotIn("Profit", header)
        self.assertIn("Revenue", header)


class FullBackupExportTests(APITestCase):
    """Admin/super only, by design — even a custom role holding every report
    permission there is (view_reports, export_reports, view_profit) must NOT get
    in, since this covers everything including user accounts and the audit log."""

    def setUp(self):
        admin_role = Role.objects.create(name="Admin", is_system_role=True)
        self.admin = User.objects.create_user(
            username="backupadmin", password="Str0ngPassw0rd!", phone="255700000023", role=admin_role, must_change_password=False
        )

        all_report_perms = Role.objects.create(name="ReportsPowerUser")
        for codename in ("view_reports", "export_reports", "view_profit"):
            perm = Permission.objects.create(codename=codename, label=codename, category="reports")
            all_report_perms.permissions.add(perm)
        self.power_user = User.objects.create_user(
            username="powerreports", password="Str0ngPassw0rd!", phone="255700000024", role=all_report_perms, must_change_password=False
        )

        category = Category.objects.create(name="Samsung")
        model = PhoneModel.objects.create(category=category, name="Galaxy A56")
        supplier = Supplier.objects.create(name="Blue Telecom")
        stock_in = StockIn.objects.create(supplier=supplier, import_date=date.today(), created_by=self.admin)
        StockItem.objects.create(
            stock_in=stock_in, category=category, model=model, quantity=2, quantity_remaining=2,
            buying_price="500000", min_selling_price="600000", max_selling_price="700000",
        )

    def test_admin_can_download_full_backup(self):
        self.client.force_authenticate(self.admin)
        res = self.client.get("/api/v1/reports/full-backup/")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertIn("spreadsheetml", res["Content-Type"])

        workbook = openpyxl.load_workbook(io.BytesIO(res.content))
        self.assertIn("Stock Items", workbook.sheetnames)
        self.assertIn("Sales", workbook.sheetnames)
        self.assertIn("Users", workbook.sheetnames)
        self.assertIn("Roles", workbook.sheetnames)
        self.assertIn("Activity Log", workbook.sheetnames)

    def test_custom_role_with_every_report_permission_is_still_denied(self):
        self.client.force_authenticate(self.power_user)
        res = self.client.get("/api/v1/reports/full-backup/")
        self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)


class LoanSalesReportTests(APITestCase):
    """The loan sales report: same customisation as the sales report, plus business and
    payment status. Fixture (revenue = price - discount; profit = revenue - buying price):

      A  Kariakoo   owner       today    2 x A56 @650k         rev 1,300k  profit 300k  paid 500k -> partial
      B  Mwanza     salesman    today    1 x Redmi @400k-20k   rev   380k  profit  80k  paid   0   -> open
      C  Kariakoo   owner       10d ago  1 x A56 @600k         rev   600k  profit 100k  paid 600k -> paid
    """

    def setUp(self):
        from loans.models import LoanPayment, LoanSale, LoanSaleItem

        perms = {
            code: Permission.objects.create(codename=code, label=code, category="x")
            for code in ("view_reports", "export_reports", "view_profit", "create_loan_sales")
        }

        def user_with(name, codes, n):
            role = Role.objects.create(name=f"role-{name}")
            role.permissions.add(*[perms[c] for c in codes])
            return User.objects.create_user(
                username=name, password="Str0ngPassw0rd!", phone=f"2557000009{n:02d}", role=role,
                must_change_password=False,
            )

        self.owner = user_with("owner", perms.keys(), 1)
        self.salesman = user_with("salesman", ("view_reports", "export_reports", "create_loan_sales"), 2)  # no view_profit
        self.reports_only = user_with("reportsonly", ("view_reports",), 3)  # no loan permission
        self.loan_only = user_with("loanonly", ("create_loan_sales",), 4)  # no view_reports

        self.samsung = Category.objects.create(name="Samsung")
        self.xiaomi = Category.objects.create(name="Xiaomi")
        a56 = PhoneModel.objects.create(category=self.samsung, name="Galaxy A56")
        redmi = PhoneModel.objects.create(category=self.xiaomi, name="Redmi 13")

        def stock(category, model, supplier_name, buying):
            batch = StockIn.objects.create(
                supplier=Supplier.objects.create(name=supplier_name), import_date=date.today(), created_by=self.owner
            )
            return StockItem.objects.create(
                stock_in=batch, category=category, model=model, quantity=20, quantity_remaining=20,
                buying_price=buying, min_selling_price="1", max_selling_price="1",
            )

        self.a56_stock = stock(self.samsung, a56, "Blue Telecom", "500000")
        self.redmi_stock = stock(self.xiaomi, redmi, "Green Mobile", "300000")

        def loan(invoice, business, seller, items, paid=0, phone="", person=""):
            record = LoanSale.objects.create(
                invoice_number=invoice, business_name=business, sold_by=seller,
                contact_person=person, contact_phone=phone, notes=f"note {invoice}",
            )
            for stock_item, price, discount in items:
                LoanSaleItem.objects.create(loan_sale=record, stock_item=stock_item, selling_price=price, discount=discount)
            if paid:
                LoanPayment.objects.create(
                    loan_sale=record, amount=paid, payment_method="cash", paid_date=date.today(), recorded_by=seller
                )
            return record

        self.loan_a = loan("LOAN-A", "Kariakoo Phones", self.owner, [(self.a56_stock, "650000", 0)] * 2, 500000, "255711", "Juma")
        self.loan_b = loan("LOAN-B", "Mwanza Mobile", self.salesman, [(self.redmi_stock, "400000", "20000")])
        self.loan_c = loan("LOAN-C", "Kariakoo Phones", self.owner, [(self.a56_stock, "600000", 0)], 600000)
        LoanSale.objects.filter(pk=self.loan_c.pk).update(created_at=timezone.now() - timedelta(days=10))

        self.client.force_authenticate(self.owner)

    URL = "/api/v1/reports/loan-sales/"

    def _wide(self, **extra):
        return {
            "date_from": (date.today() - timedelta(days=30)).isoformat(), "date_to": date.today().isoformat(), **extra,
        }

    def _today(self, **extra):
        return {"date_from": date.today().isoformat(), "date_to": date.today().isoformat(), **extra}

    def _get(self, params):
        res = self.client.get(self.URL, params)
        self.assertEqual(res.status_code, status.HTTP_200_OK, res.content)
        return res.json()

    def test_needs_both_view_reports_and_a_loan_permission(self):
        self.client.force_authenticate(self.reports_only)
        self.assertEqual(self.client.get(self.URL, self._wide()).status_code, status.HTTP_403_FORBIDDEN)
        self.client.force_authenticate(self.loan_only)
        self.assertEqual(self.client.get(self.URL, self._wide()).status_code, status.HTTP_403_FORBIDDEN)
        self.client.force_authenticate(self.salesman)
        self.assertEqual(self.client.get(self.URL, self._wide()).status_code, status.HTTP_200_OK)

    def test_grouped_by_business_shows_loans_units_revenue_profit_paid_and_owed(self):
        body = self._get(self._wide(group_by="business"))
        rows = {r["label"]: r for r in body["rows"]}
        self.assertEqual([r["label"] for r in body["rows"]], ["Kariakoo Phones", "Mwanza Mobile"])  # biggest first
        k = rows["Kariakoo Phones"]
        self.assertEqual((k["loans"], k["units"]), (2, 3))
        self.assertEqual(float(k["revenue"]), 1900000.0)
        self.assertEqual(float(k["expectedProfit"]), 400000.0)
        self.assertEqual(float(k["paid"]), 1100000.0)
        self.assertEqual(float(k["outstanding"]), 800000.0)
        m = rows["Mwanza Mobile"]
        self.assertEqual((float(m["revenue"]), float(m["paid"]), float(m["outstanding"])), (380000.0, 0.0, 380000.0))

        totals = body["totals"]
        self.assertEqual((totals["loans"], totals["units"]), (3, 4))
        self.assertEqual(float(totals["revenue"]), 2280000.0)
        self.assertEqual(float(totals["expectedProfit"]), 480000.0)
        self.assertEqual(float(totals["paid"]), 1100000.0)
        self.assertEqual(float(totals["outstanding"]), 1180000.0)

    def test_it_has_its_own_date_range_by_the_day_the_loan_was_made(self):
        today = self._get(self._today())
        self.assertEqual(today["totals"]["loans"], 2)  # C was made 10 days ago
        self.assertEqual(len(today["details"]), 2)
        self.assertEqual(self._get(self._wide())["totals"]["loans"], 3)

    def test_groupings(self):
        by_day = self._get(self._wide(group_by="day"))["rows"]
        self.assertEqual(len(by_day), 2)
        self.assertLess(by_day[0]["key"], by_day[1]["key"])  # oldest first, like the sales trend

        by_status = self._get(self._wide(group_by="status"))["rows"]
        self.assertEqual([r["label"] for r in by_status], ["Open", "Partially paid", "Paid off"])

        by_user = {r["label"]: r["loans"] for r in self._get(self._wide(group_by="user"))["rows"]}
        self.assertEqual(by_user, {"owner": 2, "salesman": 1})

    def test_grouping_by_product_has_no_payment_columns_since_payments_belong_to_the_whole_loan(self):
        for group_by, labels in (("category", {"Samsung", "Xiaomi"}), ("model", {"Galaxy A56", "Redmi 13"}),
                                 ("supplier", {"Blue Telecom", "Green Mobile"})):
            rows = self._get(self._wide(group_by=group_by))["rows"]
            self.assertEqual({r["label"] for r in rows}, labels, group_by)
            for row in rows:
                self.assertNotIn("paid", row)
                self.assertNotIn("outstanding", row)
        samsung = next(r for r in self._get(self._wide(group_by="category"))["rows"] if r["label"] == "Samsung")
        self.assertEqual((samsung["loans"], samsung["units"]), (2, 3))  # A and C both hold Samsung phones
        self.assertEqual(float(samsung["revenue"]), 1900000.0)

    def test_filters_by_status_business_and_salesperson(self):
        open_only = self._get(self._wide(status="open"))
        self.assertEqual([d["invoiceNumber"] for d in open_only["details"]], ["LOAN-B"])

        kariakoo = self._get(self._wide(business="kariakoo"))  # case-insensitive "contains"
        self.assertEqual(kariakoo["totals"]["loans"], 2)

        mine = self._get(self._wide(user=str(self.salesman.id)))
        self.assertEqual([d["invoiceNumber"] for d in mine["details"]], ["LOAN-B"])

    def test_status_filter_still_sees_the_whole_loan_when_grouped_by_product(self):
        rows = self._get(self._wide(status="partial", group_by="category"))["rows"]
        self.assertEqual([r["label"] for r in rows], ["Samsung"])
        self.assertEqual(rows[0]["units"], 2)  # loan A only

    def test_a_product_filter_counts_only_matching_items_and_drops_payment_figures(self):
        body = self._get(self._wide(category=str(self.samsung.id)))
        self.assertEqual(body["totals"]["units"], 3)
        self.assertEqual(float(body["totals"]["revenue"]), 1900000.0)
        self.assertNotIn("paid", body["totals"])
        self.assertNotIn("outstanding", body["totals"])
        for row in body["details"]:
            self.assertNotIn("paid", row)
            self.assertNotIn("balance", row)
        self.assertNotIn("paid", body["rows"][0])

    def test_detail_row_has_one_line_per_loan_with_everything(self):
        row = next(d for d in self._get(self._wide())["details"] if d["invoiceNumber"] == "LOAN-A")
        self.assertRegex(row["time"], r"^\d{2}:\d{2}$")
        self.assertEqual(row["businessName"], "Kariakoo Phones")
        self.assertEqual(row["contact"], "Juma · 255711")
        self.assertEqual(row["soldByName"], "owner")
        self.assertEqual(row["models"], "Galaxy A56 ×2")
        self.assertEqual(row["categories"], "Samsung")
        self.assertEqual(row["suppliers"], "Blue Telecom")
        self.assertEqual(row["units"], 2)
        self.assertEqual(float(row["revenue"]), 1300000.0)
        self.assertEqual(float(row["cost"]), 1000000.0)
        self.assertEqual(float(row["expectedProfit"]), 300000.0)
        self.assertEqual(float(row["paid"]), 500000.0)
        self.assertEqual(float(row["balance"]), 800000.0)
        self.assertEqual(row["status"], "Partially paid")
        self.assertEqual(row["lastPayment"], date.today().isoformat())
        self.assertEqual(row["notes"], "note LOAN-A")
        self.assertIsNone(next(d for d in self._get(self._wide())["details"] if d["invoiceNumber"] == "LOAN-B")["lastPayment"])

    def test_detail_totals_add_up_the_columns(self):
        totals = self._get(self._wide())["detailTotals"]
        self.assertEqual(totals["units"], 4)
        self.assertEqual(float(totals["revenue"]), 2280000.0)
        self.assertEqual(float(totals["paid"]), 1100000.0)
        self.assertEqual(float(totals["balance"]), 1180000.0)

    def test_cost_and_profit_are_hidden_without_view_profit(self):
        self.client.force_authenticate(self.salesman)
        body = self._get(self._wide())
        self.assertNotIn("expectedProfit", body["totals"])
        for row in body["rows"]:
            self.assertNotIn("expectedProfit", row)
        for row in body["details"]:
            self.assertNotIn("cost", row)
            self.assertNotIn("expectedProfit", row)
        self.assertNotIn("expectedProfit", body["detailTotals"])
        self.assertIn("revenue", body["totals"])  # revenue and payments stay visible

    def test_an_overpaid_loan_has_zero_balance_and_cannot_cancel_another_loans_debt(self):
        from loans.models import LoanPayment

        LoanPayment.objects.create(
            loan_sale=self.loan_b, amount="500000", payment_method="cash", paid_date=date.today(), recorded_by=self.owner
        )  # B is worth 380k
        body = self._get(self._wide(group_by="status"))
        paid_off = next(r for r in body["rows"] if r["label"] == "Paid off")
        self.assertEqual(paid_off["loans"], 2)
        self.assertEqual(float(paid_off["outstanding"]), 0.0)
        self.assertEqual(float(body["totals"]["outstanding"]), 800000.0)  # only A still owes

    def test_bad_group_by_or_status_falls_back_instead_of_erroring(self):
        body = self._get(self._wide(group_by="nonsense", status="whatever"))
        self.assertEqual({r["label"] for r in body["rows"]}, {"Kariakoo Phones", "Mwanza Mobile"})
        self.assertEqual(body["totals"]["loans"], 3)

    def test_an_empty_range_is_an_empty_report(self):
        body = self._get({"date_from": "2001-01-01", "date_to": "2001-01-02"})
        self.assertEqual((body["rows"], body["details"]), ([], []))
        self.assertEqual(body["totals"]["loans"], 0)

    def test_xlsx_export_has_a_summary_and_details_sheet(self):
        res = self.client.get(self.URL, {**self._wide(), "export": "xlsx"})
        workbook = openpyxl.load_workbook(io.BytesIO(res.content))
        self.assertEqual(workbook.sheetnames, ["Summary", "Details"])
        summary_header = [c.value for c in workbook["Summary"][1]]
        self.assertEqual(summary_header, ["Business", "Loans", "Units", "Revenue", "Expected Profit", "Paid", "Still Owed"])
        details = list(workbook["Details"].iter_rows(values_only=True))
        self.assertIn("Balance", details[0])
        self.assertEqual(details[-1][0], "Total")

    def test_exports_leave_out_profit_without_view_profit_and_the_pdf_renders(self):
        self.client.force_authenticate(self.salesman)
        res = self.client.get(self.URL, {**self._wide(), "export": "xlsx"})
        workbook = openpyxl.load_workbook(io.BytesIO(res.content))
        self.assertNotIn("Expected Profit", [c.value for c in workbook["Summary"][1]])
        self.assertNotIn("Cost", [c.value for c in workbook["Details"][1]])
        pdf = self.client.get(self.URL, {**self._wide(), "export": "pdf"})
        self.assertEqual(pdf.status_code, status.HTTP_200_OK)
        self.assertTrue(pdf.content.startswith(b"%PDF"))

    def test_exporting_needs_export_reports(self):
        role = self.loan_only.role
        role.permissions.add(Permission.objects.get(codename="view_reports"))
        self.client.force_authenticate(self.loan_only)  # can view, but may not export
        self.assertEqual(self.client.get(self.URL, self._wide()).status_code, status.HTTP_200_OK)
        self.assertEqual(self.client.get(self.URL, {**self._wide(), "export": "xlsx"}).status_code, 403)
