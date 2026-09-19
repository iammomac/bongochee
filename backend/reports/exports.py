"""Generic export layer — every report view produces the same (headers, rows) shape,
so Excel/PDF generation is written once here rather than per report type."""

import io
from decimal import Decimal
from xml.sax.saxutils import escape

import openpyxl
from django.http import HttpResponse
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

BRAND_PRIMARY = colors.HexColor("#25B1FF")
ROW_ALT = colors.HexColor("#F7F8FC")
GRID = colors.HexColor("#E5E7EB")


def _write_sheet(sheet, headers, rows, total_row=None):
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    if total_row:
        sheet.append(list(total_row))
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
    for column_cells in sheet.columns:
        width = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(width + 2, 10), 40)


def rows_to_xlsx(filename, headers, rows, detail=None):
    """detail: optional {"headers", "rows", "total_row"} -- written as a second
    "Details" sheet so the workbook carries both the grouped summary and every
    underlying transaction."""
    workbook = openpyxl.Workbook()
    summary_sheet = workbook.active
    _write_sheet(summary_sheet, headers, rows)
    if detail:
        summary_sheet.title = "Summary"
        _write_sheet(workbook.create_sheet(title="Details"), detail["headers"], detail["rows"], detail.get("total_row"))

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    workbook.save(response)
    return response


def multi_sheet_xlsx(filename, sheets):
    """sheets: list of (sheet_name, headers, rows) — one tab per entry, in order."""
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for sheet_name, headers, rows in sheets:
        sheet = workbook.create_sheet(title=sheet_name[:31])  # Excel's own sheet-name limit
        _write_sheet(sheet, headers, rows)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    workbook.save(response)
    return response


PDF_MARGIN = 24


def _pdf_text(cell):
    if cell is None:
        return ""
    if isinstance(cell, bool):
        return "Yes" if cell else "No"
    if isinstance(cell, (Decimal, float)):
        return f"{cell:,.0f}"
    if isinstance(cell, int):
        return f"{cell:,}"
    return str(cell)


def _column_widths(headers, rows, total_width):
    """Widths proportional to how much each column's content needs, capped so a long
    note wraps onto extra lines instead of squeezing every other column."""
    weights = []
    for index, header in enumerate(headers):
        # Headers are bold (wider per character); +3 leaves room for the cell padding
        # so a short word like "Samsung" or "Discount" never breaks mid-word.
        header_len = int(len(str(header)) * 1.2) + 1
        longest = max([header_len] + [len(_pdf_text(row[index])) for row in rows[:200]])
        # Digits and capitals run wider than the average letter, hence the 10% margin.
        weights.append(min(max(int(longest * 1.1), 4), 28) + 3)
    scale = total_width / sum(weights)
    return [weight * scale for weight in weights]


def _pdf_table(headers, rows, font_size, total_row=None):
    body = ParagraphStyle("cell", fontName="Helvetica", fontSize=font_size, leading=font_size + 2)
    bold = ParagraphStyle("cell-bold", parent=body, fontName="Helvetica-Bold")
    head = ParagraphStyle("cell-head", parent=bold, textColor=colors.white)

    def line(cells, style):
        return [Paragraph(escape(_pdf_text(cell)), style) for cell in cells]

    all_rows = list(rows) + ([total_row] if total_row else [])
    data = [line(headers, head)] + [line(row, body) for row in rows]
    if total_row:
        data.append(line(total_row, bold))

    page_width = landscape(A4)[0] - 2 * PDF_MARGIN
    table = Table(data, colWidths=_column_widths(headers, all_rows, page_width), repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), BRAND_PRIMARY),
                ("GRID", (0, 0), (-1, -1), 0.5, GRID),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, ROW_ALT]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]
        )
    )
    return table


def rows_to_pdf(filename, title, headers, rows, detail=None):
    """detail: optional {"headers", "rows", "total_row"} -- rendered as a second,
    smaller-print table under the summary."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=PDF_MARGIN,
        rightMargin=PDF_MARGIN,
        topMargin=PDF_MARGIN,
        bottomMargin=PDF_MARGIN,
    )
    styles = getSampleStyleSheet()
    elements = [Paragraph(title, styles["Title"]), Spacer(1, 12)]

    # The wide reports (e.g. Loss, which is itself a per-sale detail table) need the
    # same small print as the detail section to fit a landscape page.
    elements.append(_pdf_table(headers, rows, font_size=9 if len(headers) <= 6 else 6.5))
    if detail:
        elements.append(Spacer(1, 18))
        elements.append(Paragraph("Detailed transactions", styles["Heading2"]))
        elements.append(Spacer(1, 6))
        elements.append(_pdf_table(detail["headers"], detail["rows"], font_size=6.5, total_row=detail.get("total_row")))
    doc.build(elements)

    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}.pdf"'
    return response
