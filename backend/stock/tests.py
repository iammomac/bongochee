import io
from datetime import date

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework import status
from rest_framework.test import APITestCase

from accounts.models import User
from catalog.models import Category, PhoneModel
from rbac.models import Permission, Role
from sales.models import Sale, SaleItem
from stock.models import StockIn, StockItem
from suppliers.models import Supplier

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def build_xlsx(header, rows):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


class StockInTests(APITestCase):
    def setUp(self):
        perm = Permission.objects.create(codename="add_stock", label="Add Stock", category="stock")
        role = Role.objects.create(name="Stocker")
        role.permissions.add(perm)
        self.user = User.objects.create_user(
            username="stocker",
            password="Str0ngPassw0rd!",
            phone="255700000030",
            role=role,
            must_change_password=False,
        )
        self.client.force_authenticate(self.user)
        self.supplier = Supplier.objects.create(name="Blue Telecom")

    def test_stock_in_requires_at_least_one_item(self):
        res = self.client.post(
            "/api/v1/stock/stock-ins/",
            {"supplier": str(self.supplier.id), "importDate": str(date.today()), "items": []},
            format="json",
        )
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)

    def test_stock_in_creates_nested_items(self):
        category = Category.objects.create(name="Samsung")
        model = PhoneModel.objects.create(category=category, name="Galaxy A56")
        res = self.client.post(
            "/api/v1/stock/stock-ins/",
            {
                "supplier": str(self.supplier.id),
                "importDate": str(date.today()),
                "invoiceNumber": "INV-100",
                "items": [
                    {
                        "category": str(category.id),
                        "model": str(model.id),
                        "quantity": 10,
                        "buyingPrice": "500000",
                        "minSellingPrice": "600000",
                        "maxSellingPrice": "700000",
                    }
                ],
            },
            format="json",
        )
        self.assertEqual(res.status_code, status.HTTP_201_CREATED)
        self.assertEqual(StockItem.objects.count(), 1)
        item = StockItem.objects.first()
        self.assertEqual(item.quantity_remaining, 10)

    def test_import_preview_resolves_and_creates_idempotently(self):
        header = ["Category", "Model", "Quantity", "Buying Price", "Min Selling Price", "Max Selling Price"]
        file_one = build_xlsx(header, [("Samsung", "Galaxy A56", 10, 500000, 600000, 700000)])
        upload = SimpleUploadedFile("stock.xlsx", file_one.read(), content_type=XLSX_CONTENT_TYPE)
        res = self.client.post("/api/v1/stock/import-preview/", {"file": upload}, format="multipart")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        rows = res.json()["rows"]
        self.assertEqual(len(rows), 1)
        self.assertIsNone(rows[0]["error"])
        self.assertEqual(rows[0]["category"]["name"], "Samsung")
        self.assertEqual(Category.objects.filter(name="Samsung").count(), 1)

        file_two = build_xlsx(header, [("samsung", "galaxy a56", 5, 500000, 600000, 700000)])
        upload_two = SimpleUploadedFile("stock2.xlsx", file_two.read(), content_type=XLSX_CONTENT_TYPE)
        self.client.post("/api/v1/stock/import-preview/", {"file": upload_two}, format="multipart")
        self.assertEqual(Category.objects.filter(name__iexact="samsung").count(), 1)
        self.assertEqual(PhoneModel.objects.filter(name__iexact="galaxy a56").count(), 1)

    def test_import_preview_reports_row_errors_without_dropping_them(self):
        header = ["Category", "Model", "Quantity", "Buying Price", "Min Selling Price", "Max Selling Price"]
        file = build_xlsx(header, [("Samsung", "Galaxy A56", "not-a-number", 500000, 600000, 700000)])
        upload = SimpleUploadedFile("stock.xlsx", file.read(), content_type=XLSX_CONTENT_TYPE)
        res = self.client.post("/api/v1/stock/import-preview/", {"file": upload}, format="multipart")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        rows = res.json()["rows"]
        self.assertEqual(len(rows), 1)
        self.assertIn("Quantity", rows[0]["error"])

    def test_import_preview_requires_expected_columns(self):
        file = build_xlsx(["Category", "Model"], [("Samsung", "Galaxy A56")])
        upload = SimpleUploadedFile("bad.xlsx", file.read(), content_type=XLSX_CONTENT_TYPE)
        res = self.client.post("/api/v1/stock/import-preview/", {"file": upload}, format="multipart")
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)


class StockItemAccessTests(APITestCase):
    def setUp(self):
        create_sales_perm = Permission.objects.create(codename="create_sales", label="Create Sales", category="sales")
        role = Role.objects.create(name="Seller")
        role.permissions.add(create_sales_perm)
        self.seller = User.objects.create_user(
            username="seller3",
            password="Str0ngPassw0rd!",
            phone="255700000050",
            role=role,
            must_change_password=False,
        )
        self.client.force_authenticate(self.seller)

        supplier = Supplier.objects.create(name="Blue Telecom")
        category = Category.objects.create(name="Samsung")
        model = PhoneModel.objects.create(category=category, name="Galaxy A56")
        stock_in = StockIn.objects.create(supplier=supplier, import_date=date.today(), created_by=self.seller)
        self.in_stock = StockItem.objects.create(
            stock_in=stock_in,
            category=category,
            model=model,
            quantity=5,
            quantity_remaining=5,
            buying_price="500000",
            min_selling_price="600000",
            max_selling_price="700000",
        )
        self.sold_out = StockItem.objects.create(
            stock_in=stock_in,
            category=category,
            model=model,
            quantity=3,
            quantity_remaining=0,
            buying_price="500000",
            min_selling_price="600000",
            max_selling_price="700000",
        )

    def test_create_sales_only_user_can_list_stock_items(self):
        res = self.client.get("/api/v1/stock/stock-items/")
        self.assertEqual(res.status_code, status.HTTP_200_OK)

    def test_quantity_remaining_filter_excludes_sold_out_batches(self):
        res = self.client.get("/api/v1/stock/stock-items/", {"quantity_remaining__gt": 0})
        ids = [item["id"] for item in res.json()["results"]] if "results" in res.json() else [
            item["id"] for item in res.json()
        ]
        self.assertIn(str(self.in_stock.id), ids)
        self.assertNotIn(str(self.sold_out.id), ids)


class StockItemDeleteTests(APITestCase):
    def setUp(self):
        add_perm = Permission.objects.create(codename="add_stock", label="Add Stock", category="stock")
        delete_perm = Permission.objects.create(codename="delete_stock", label="Delete Stock", category="stock")
        role = Role.objects.create(name="Stocker")
        role.permissions.add(add_perm, delete_perm)
        self.stocker = User.objects.create_user(
            username="stocker4", password="Str0ngPassw0rd!", phone="255700000053", role=role, must_change_password=False
        )

        add_only_role = Role.objects.create(name="StockerNoDelete")
        add_only_role.permissions.add(add_perm)
        self.stocker_no_delete = User.objects.create_user(
            username="stocker5", password="Str0ngPassw0rd!", phone="255700000054", role=add_only_role, must_change_password=False
        )

        supplier = Supplier.objects.create(name="Blue Telecom")
        category = Category.objects.create(name="Samsung")
        model = PhoneModel.objects.create(category=category, name="Galaxy A56")
        self.stock_in = StockIn.objects.create(supplier=supplier, import_date=date.today(), created_by=self.stocker)
        self.untouched = StockItem.objects.create(
            stock_in=self.stock_in, category=category, model=model,
            quantity=5, quantity_remaining=5,
            buying_price="500000", min_selling_price="600000", max_selling_price="700000",
        )
        self.partially_sold = StockItem.objects.create(
            stock_in=self.stock_in, category=category, model=model,
            quantity=5, quantity_remaining=4,
            buying_price="500000", min_selling_price="600000", max_selling_price="700000",
        )
        sale = Sale.objects.create(invoice_number="INV-DEL-1", customer_name="X", payment_method="cash", sold_by=self.stocker)
        SaleItem.objects.create(sale=sale, stock_item=self.partially_sold, selling_price="650000")

    def test_untouched_line_can_be_deleted(self):
        self.client.force_authenticate(self.stocker)
        res = self.client.delete(f"/api/v1/stock/stock-items/{self.untouched.id}/")
        self.assertEqual(res.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(StockItem.objects.filter(id=self.untouched.id).exists())

    def test_line_with_a_sale_against_it_cannot_be_deleted(self):
        self.client.force_authenticate(self.stocker)
        res = self.client.delete(f"/api/v1/stock/stock-items/{self.partially_sold.id}/")
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertTrue(StockItem.objects.filter(id=self.partially_sold.id).exists())

    def test_user_without_delete_stock_permission_is_denied(self):
        self.client.force_authenticate(self.stocker_no_delete)
        res = self.client.delete(f"/api/v1/stock/stock-items/{self.untouched.id}/")
        self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)
        self.assertTrue(StockItem.objects.filter(id=self.untouched.id).exists())


class StockItemNotesTests(APITestCase):
    def setUp(self):
        perm = Permission.objects.create(codename="add_stock", label="Add Stock", category="stock")
        role = Role.objects.create(name="Stocker")
        role.permissions.add(perm)
        self.user = User.objects.create_user(
            username="stocker6", password="Str0ngPassw0rd!", phone="255700000055", role=role, must_change_password=False
        )
        self.client.force_authenticate(self.user)
        self.supplier = Supplier.objects.create(name="Blue Telecom")

    def test_notes_saved_and_returned_per_line(self):
        category = Category.objects.create(name="Samsung")
        model = PhoneModel.objects.create(category=category, name="Galaxy A56")
        res = self.client.post(
            "/api/v1/stock/stock-ins/",
            {
                "supplier": str(self.supplier.id),
                "importDate": str(date.today()),
                "items": [
                    {
                        "category": str(category.id), "model": str(model.id),
                        "quantity": 2, "buyingPrice": "500000",
                        "minSellingPrice": "600000", "maxSellingPrice": "700000",
                        "notes": "Full box, unused",
                    }
                ],
            },
            format="json",
        )
        self.assertEqual(res.status_code, status.HTTP_201_CREATED)
        item = StockItem.objects.get()
        self.assertEqual(item.notes, "Full box, unused")
        list_res = self.client.get("/api/v1/stock/stock-items/")
        body = list_res.json()
        rows = body["results"] if "results" in body else body
        self.assertEqual(rows[0]["notes"], "Full box, unused")


class StockImportTemplateTests(APITestCase):
    def setUp(self):
        add_stock_perm = Permission.objects.create(codename="add_stock", label="Add Stock", category="stock")
        stocker_role = Role.objects.create(name="Stocker")
        stocker_role.permissions.add(add_stock_perm)
        self.stocker = User.objects.create_user(
            username="stocker3", password="Str0ngPassw0rd!", phone="255700000051", role=stocker_role, must_change_password=False
        )

        no_perm_role = Role.objects.create(name="NoPerms")
        self.outsider = User.objects.create_user(
            username="outsider2", password="Str0ngPassw0rd!", phone="255700000052", role=no_perm_role, must_change_password=False
        )

    def test_add_stock_holder_can_download_template(self):
        self.client.force_authenticate(self.stocker)
        res = self.client.get("/api/v1/stock/import-template/")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(
            res["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        workbook = openpyxl.load_workbook(io.BytesIO(res.content))
        self.assertEqual(workbook.sheetnames, ["Stock Import Template", "Instructions"])
        header_row = [cell.value for cell in workbook["Stock Import Template"][1]]
        self.assertEqual(
            header_row, ["Category", "Model", "Quantity", "Buying Price", "Min Selling Price", "Max Selling Price"]
        )

    def test_user_without_add_stock_is_denied(self):
        self.client.force_authenticate(self.outsider)
        res = self.client.get("/api/v1/stock/import-template/")
        self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)


class StockItemEditTests(APITestCase):
    """Editing a stock line as a role that only has Edit Stock (no Add Stock)."""

    def setUp(self):
        edit_perm = Permission.objects.create(codename="edit_stock", label="Edit Stock", category="stock")
        sales_perm = Permission.objects.create(codename="create_sales", label="Create Sales", category="sales")
        editor_role = Role.objects.create(name="Stock editor")
        editor_role.permissions.add(edit_perm)
        seller_role = Role.objects.create(name="Seller")
        seller_role.permissions.add(sales_perm)
        self.editor = User.objects.create_user(
            username="editor", password="Str0ngPassw0rd!", phone="255700000070", role=editor_role, must_change_password=False
        )
        self.seller = User.objects.create_user(
            username="seller9", password="Str0ngPassw0rd!", phone="255700000071", role=seller_role, must_change_password=False
        )

        self.supplier = Supplier.objects.create(name="Blue Telecom")
        self.other_supplier = Supplier.objects.create(name="Red Mobile")
        self.samsung = Category.objects.create(name="Samsung")
        self.apple = Category.objects.create(name="Apple")
        self.a56 = PhoneModel.objects.create(category=self.samsung, name="Galaxy A56")
        self.s23 = PhoneModel.objects.create(category=self.samsung, name="Galaxy S23")
        self.iphone = PhoneModel.objects.create(category=self.apple, name="iPhone 15")
        stock_in = StockIn.objects.create(
            supplier=self.supplier, import_date=date(2026, 9, 1), invoice_number="INV-1", created_by=self.editor
        )
        # 10 received, 4 sold, 6 still in stock.
        self.item = self._line(stock_in, self.a56, quantity=10, remaining=6)
        self.sibling = self._line(stock_in, self.s23, quantity=3, remaining=3)
        self.client.force_authenticate(self.editor)

    def _line(self, stock_in, model, quantity, remaining):
        return StockItem.objects.create(
            stock_in=stock_in,
            category=model.category,
            model=model,
            quantity=quantity,
            quantity_remaining=remaining,
            buying_price="500000",
            min_selling_price="600000",
            max_selling_price="700000",
        )

    def _patch(self, payload, item=None):
        return self.client.patch(f"/api/v1/stock/stock-items/{(item or self.item).id}/", payload, format="json")

    def test_editor_without_add_stock_can_open_the_list_and_edit(self):
        self.assertEqual(self.client.get("/api/v1/stock/stock-items/").status_code, status.HTTP_200_OK)
        res = self._patch({"notes": "Full box"})
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.item.refresh_from_db()
        self.assertEqual(self.item.notes, "Full box")

    def test_a_role_without_edit_stock_cannot_edit(self):
        self.client.force_authenticate(self.seller)
        self.assertEqual(self._patch({"notes": "x"}).status_code, status.HTTP_403_FORBIDDEN)

    def test_raising_the_quantity_adds_the_difference_to_what_is_in_stock(self):
        res = self._patch({"quantity": 15})
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.item.refresh_from_db()
        self.assertEqual(self.item.quantity, 15)
        self.assertEqual(self.item.quantity_remaining, 11)  # 4 sold stay sold
        self.assertEqual(res.json()["quantityRemaining"], 11)

    def test_lowering_the_quantity_takes_the_difference_out_of_stock(self):
        self._patch({"quantity": 7})
        self.item.refresh_from_db()
        self.assertEqual((self.item.quantity, self.item.quantity_remaining), (7, 3))

    def test_quantity_can_go_down_to_exactly_what_was_sold(self):
        res = self._patch({"quantity": 4})
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.item.refresh_from_db()
        self.assertEqual(self.item.quantity_remaining, 0)

    def test_quantity_below_what_was_sold_is_refused_and_nothing_changes(self):
        res = self._patch({"quantity": 3, "notes": "should not stick"})
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("4 already sold", res.json()["detail"])
        self.item.refresh_from_db()
        self.assertEqual((self.item.quantity, self.item.quantity_remaining, self.item.notes), (10, 6, ""))

    def test_a_model_from_another_brand_is_refused(self):
        res = self._patch({"model": str(self.iphone.id)})  # still under Samsung
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("isn't a Samsung model", res.json()["detail"])

    def test_changing_brand_and_model_together_works(self):
        res = self._patch({"category": str(self.apple.id), "model": str(self.iphone.id)})
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        body = res.json()
        self.assertEqual((body["categoryName"], body["modelName"]), ("Apple", "iPhone 15"))

    def test_max_price_below_min_is_refused(self):
        res = self._patch({"maxSellingPrice": "550000"})
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("max selling price", res.json()["detail"])

    def test_prices_and_notes_update(self):
        res = self._patch({"buyingPrice": "480000", "minSellingPrice": "590000", "maxSellingPrice": "690000", "notes": "used"})
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.item.refresh_from_db()
        self.assertEqual(str(self.item.buying_price), "480000.00")
        self.assertEqual(str(self.item.min_selling_price), "590000.00")
        self.assertEqual(self.item.notes, "used")

    def test_supplier_date_and_invoice_are_corrected_for_the_whole_batch(self):
        res = self._patch({"supplier": str(self.other_supplier.id), "importDate": "2026-08-30", "invoiceNumber": "INV-9"})
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        body = res.json()
        self.assertEqual((body["supplierName"], body["importDate"], body["invoiceNumber"]), ("Red Mobile", "2026-08-30", "INV-9"))
        self.assertEqual(body["batchSize"], 2)
        self.sibling.refresh_from_db()
        self.assertEqual(self.sibling.stock_in.supplier, self.other_supplier)
        self.assertEqual(self.sibling.stock_in.invoice_number, "INV-9")

    def test_list_reports_the_batch_details_the_edit_form_needs(self):
        body = self.client.get("/api/v1/stock/stock-items/").json()
        row = next(r for r in body["results"] if r["id"] == str(self.item.id))
        self.assertEqual(row["supplier"], str(self.supplier.id))
        self.assertEqual(row["invoiceNumber"], "INV-1")
        self.assertEqual(row["batchSize"], 2)

    def test_an_edit_is_written_to_the_activity_log(self):
        from activitylog.models import ActivityLog

        self._patch({"quantity": 12, "notes": "boxed"})
        entry = ActivityLog.objects.get(action="stock.update")
        self.assertEqual(entry.user, self.editor)
        self.assertEqual(entry.details["changes"]["quantity"], [10, 12])
        self.assertEqual(entry.details["changes"]["in_stock"], [6, 8])
        self.assertEqual(entry.details["changes"]["notes"], ["", "boxed"])
        self.assertNotIn("buying_price", entry.details["changes"])

    def test_an_edit_that_changes_nothing_is_not_logged(self):
        from activitylog.models import ActivityLog

        self._patch({"quantity": 10, "notes": ""})
        self.assertFalse(ActivityLog.objects.filter(action="stock.update").exists())

    def test_search_finds_a_line_by_supplier_or_invoice(self):
        by_supplier = self.client.get("/api/v1/stock/stock-items/", {"search": "Blue"}).json()["results"]
        self.assertEqual(len(by_supplier), 2)
        by_invoice = self.client.get("/api/v1/stock/stock-items/", {"search": "INV-1"}).json()["results"]
        self.assertEqual(len(by_invoice), 2)
        self.assertEqual(self.client.get("/api/v1/stock/stock-items/", {"search": "Nokia"}).json()["results"], [])

    def test_pagination_reports_the_total_so_older_lines_can_be_reached(self):
        stock_in = StockIn.objects.create(supplier=self.supplier, import_date=date(2026, 9, 2), created_by=self.editor)
        for _ in range(30):
            self._line(stock_in, self.a56, quantity=1, remaining=1)
        first = self.client.get("/api/v1/stock/stock-items/").json()
        second = self.client.get("/api/v1/stock/stock-items/", {"page": 2}).json()
        self.assertEqual(first["count"], 32)
        self.assertEqual(len(first["results"]), 25)
        self.assertEqual(len(second["results"]), 7)
        self.assertFalse({r["id"] for r in first["results"]} & {r["id"] for r in second["results"]})


class StockSearchTests(APITestCase):
    """?search= on the stock list matches anything shown in the table."""

    def setUp(self):
        perm = Permission.objects.create(codename="edit_stock", label="Edit Stock", category="stock")
        role = Role.objects.create(name="Stock editor")
        role.permissions.add(perm)
        self.user = User.objects.create_user(
            username="searcher", password="Str0ngPassw0rd!", phone="255700000077", role=role, must_change_password=False
        )
        self.client.force_authenticate(self.user)

        blue = Supplier.objects.create(name="Blue Telecom")
        red = Supplier.objects.create(name="Red Mobile")
        samsung = Category.objects.create(name="Samsung")
        apple = Category.objects.create(name="Apple")

        def line(supplier, brand, model_name, quantity, remaining, buying, imported, invoice="", notes=""):
            model, _ = PhoneModel.objects.get_or_create(category=brand, name=model_name)
            batch = StockIn.objects.create(
                supplier=supplier, import_date=imported, invoice_number=invoice, created_by=self.user
            )
            return StockItem.objects.create(
                stock_in=batch,
                category=brand,
                model=model,
                quantity=quantity,
                quantity_remaining=remaining,
                buying_price=buying,
                min_selling_price=buying + 100000,
                max_selling_price=buying + 200000,
                notes=notes,
            )

        self.s23 = line(blue, samsung, "Galaxy S23", 20, 12, 900000, date(2026, 9, 3), "INV-77", "full box")
        self.a56 = line(red, samsung, "Galaxy A56", 10, 3, 500000, date(2026, 8, 15), "INV-88")
        self.iphone = line(blue, apple, "iPhone 15", 8, 0, 1750000, date(2026, 9, 20), "INV-99", "used")

    def _found(self, text):
        res = self.client.get("/api/v1/stock/stock-items/", {"search": text})
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        return {row["modelName"] for row in res.json()["results"]}

    def test_by_model_brand_supplier_invoice_and_notes(self):
        self.assertEqual(self._found("s23"), {"Galaxy S23"})
        self.assertEqual(self._found("apple"), {"iPhone 15"})
        self.assertEqual(self._found("red mobile"), {"Galaxy A56"})
        self.assertEqual(self._found("inv-99"), {"iPhone 15"})
        self.assertEqual(self._found("full box"), {"Galaxy S23"})

    def test_every_word_must_match_but_each_can_match_a_different_thing(self):
        self.assertEqual(self._found("galaxy blue"), {"Galaxy S23"})  # model + supplier
        self.assertEqual(self._found("samsung used"), set())  # no Samsung line is "used"

    def test_by_quantity_received_or_in_stock(self):
        self.assertEqual(self._found("20"), {"Galaxy S23"})  # received
        self.assertEqual(self._found("12"), {"Galaxy S23"})  # in stock
        # 3 is the A56's quantity in stock, and also part of the S23's model name.
        self.assertEqual(self._found("3"), {"Galaxy A56", "Galaxy S23"})
        self.assertEqual(self._found("received 10"), {"Galaxy A56"})  # the label word is ignored

    def test_a_short_number_is_a_quantity_not_a_price_fragment(self):
        zed = StockItem.objects.create(
            stock_in=StockIn.objects.create(supplier=Supplier.objects.first(), import_date=date(2026, 12, 25), created_by=self.user),
            category=Category.objects.get(name="Samsung"),
            model=PhoneModel.objects.create(category=Category.objects.get(name="Samsung"), name="Zed"),
            quantity=40,
            quantity_remaining=40,
            buying_price=300000,
            min_selling_price=400000,
            max_selling_price=500000,
        )
        # "3" sits inside 300000.00 but a short number isn't matched against prices...
        self.assertNotIn("Zed", self._found("3"))
        # ...while the full price is, and so is the exact quantity.
        self.assertIn("Zed", self._found("300000"))
        self.assertEqual(self._found("40"), {"Zed"})
        self.assertEqual(zed.quantity, 40)

    def test_by_price(self):
        self.assertEqual(self._found("900000"), {"Galaxy S23"})  # buying price
        self.assertEqual(self._found("1,750,000"), {"iPhone 15"})
        self.assertEqual(self._found("600000"), {"Galaxy A56"})  # min selling price
        self.assertEqual(self._found("1000000"), {"Galaxy S23"})  # 900000 + 100000 min selling

    def test_by_date_year_month_or_exact_day(self):
        self.assertEqual(self._found("2026-08"), {"Galaxy A56"})
        self.assertEqual(self._found("2026-09-03"), {"Galaxy S23"})
        self.assertEqual(self._found("03/09/2026"), {"Galaxy S23"})
        self.assertEqual(self._found("2026"), {"Galaxy S23", "Galaxy A56", "iPhone 15"})
        self.assertEqual(self._found("sep"), {"Galaxy S23", "iPhone 15"})
        self.assertEqual(self._found("august"), {"Galaxy A56"})

    def test_by_status(self):
        self.assertEqual(self._found("out of stock"), {"iPhone 15"})
        self.assertEqual(self._found("sold out"), {"iPhone 15"})
        self.assertEqual(self._found("low"), {"Galaxy A56"})  # 3 left
        self.assertEqual(self._found("low stock"), {"Galaxy A56"})
        self.assertEqual(self._found("healthy"), {"Galaxy S23"})
        self.assertEqual(self._found("in stock"), {"Galaxy S23", "Galaxy A56"})

    def test_status_words_combine_with_other_words(self):
        self.assertEqual(self._found("samsung low"), {"Galaxy A56"})
        self.assertEqual(self._found("blue out of stock"), {"iPhone 15"})

    def test_search_is_case_insensitive_and_blank_shows_everything(self):
        self.assertEqual(self._found("GALAXY S23"), {"Galaxy S23"})
        self.assertEqual(len(self._found("   ")), 3)
        self.assertEqual(self._found("nothing like this"), set())

    def test_search_still_works_with_the_sales_style_filter(self):
        res = self.client.get("/api/v1/stock/stock-items/", {"search": "iphone", "quantity_remaining__gt": 0})
        self.assertEqual(res.json()["results"], [])
